import os
from typing import List, Dict, Any
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cleaner import sanitize_listing


COLUMN_MAPPING = {
    "name": "Business Name",
    "phone": "Phone Number",
    "whatsapp": "WhatsApp Number",
    "email": "Email ID",
    "contact_person": "Contact Person",
    "full_address": "Full Address",
    "area": "Area / Locality",
    "city": "City",
    "pincode": "Pincode",
    "rating": "Rating",
    "total_reviews": "Total Reviews",
    "website": "Website",
    "justdial_url": "Source URL"
}


def prepare_cleaned_rows(listings: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Sanitizes each listing and maps to standardized column headers."""
    rows = []
    for item in listings:
        clean_item = sanitize_listing(item)
        row = {}
        for key, col_name in COLUMN_MAPPING.items():
            val = clean_item.get(key, "")
            row[col_name] = val
        rows.append(row)
    return rows


def export_to_excel(listings: List[Dict[str, Any]], output_filepath: str) -> str:
    """
    Exports scraped listings into a professionally styled Excel file (.xlsx).
    Cleans data, preserves leading zeros for phone numbers, and formats headers, borders, and widths.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    rows = prepare_cleaned_rows(listings)
    df = pd.DataFrame(rows, columns=list(COLUMN_MAPPING.values()))

    with pd.ExcelWriter(output_filepath, engine="openpyxl") as writer:
        sheet_name = "Leads"
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.sheets[sheet_name]

        # Freeze panes on header row
        worksheet.freeze_panes = "A2"

        # Enable auto-filter
        worksheet.auto_filter.ref = worksheet.dimensions

        # Styling definitions
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy blue
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side
        )

        data_font = Font(name="Calibri", size=10, color="1E293B")
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")
        data_align_right = Alignment(horizontal="right", vertical="center")

        # Format Header Row
        worksheet.row_dimensions[1].height = 28
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = cell_border

        # Columns formatting
        text_cols = {"Phone Number", "WhatsApp Number", "Pincode"}
        numeric_cols = {"Rating", "Total Reviews"}
        center_cols = {"Rating", "Total Reviews", "Pincode", "Phone Number", "WhatsApp Number"}

        for row_idx in range(2, len(df) + 2):
            worksheet.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)
            row_fill = white_fill if is_even else zebra_fill

            for col_idx, col_name in enumerate(df.columns, start=1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.fill = row_fill
                cell.border = cell_border

                raw_val = cell.value

                # Convert numeric columns to actual numbers for Excel sorting
                if col_name == "Rating" and raw_val:
                    try:
                        cell.value = float(raw_val)
                        cell.number_format = "0.0"
                    except Exception:
                        pass
                elif col_name == "Total Reviews" and raw_val:
                    try:
                        cell.value = int(str(raw_val).replace(",", ""))
                        cell.number_format = "#,##0"
                    except Exception:
                        pass

                if col_name in center_cols:
                    cell.alignment = data_align_center
                elif col_name in numeric_cols:
                    cell.alignment = data_align_right
                else:
                    cell.alignment = data_align_left

                # Force text format so phone numbers don't lose leading zeros
                if col_name in text_cols:
                    cell.number_format = "@"

        # Auto-adjust column widths
        for col_idx, col_name in enumerate(df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(col_name)
            for row in range(2, min(len(df) + 2, 120)):
                val_str = str(worksheet.cell(row=row, column=col_idx).value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)

            adjusted_width = min(max(max_len + 4, 12), 55)
            worksheet.column_dimensions[col_letter].width = adjusted_width

    return output_filepath


def export_to_csv(listings: List[Dict[str, Any]], output_filepath: str) -> str:
    """Exports clean scraped listings to CSV."""
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    rows = prepare_cleaned_rows(listings)
    df = pd.DataFrame(rows, columns=list(COLUMN_MAPPING.values()))
    df.to_csv(output_filepath, index=False, encoding="utf-8-sig")
    return output_filepath
