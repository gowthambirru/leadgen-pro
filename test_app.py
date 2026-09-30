import time
import os
import openpyxl
from fastapi.testclient import TestClient
from app import app, jobs

client = TestClient(app)

def test_full_pipeline():
    print("[1] Testing GET / index page...")
    resp = client.get("/")
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
    assert "Universal Lead" in resp.text or "Extractor" in resp.text
    print(" -> Index page served successfully.")

    # ----------------------------------------------------
    # Test 1: Direct Justdial URL Scraper
    # ----------------------------------------------------
    print("\n[2] Testing POST /api/scrape (Direct Justdial)...")
    target_url = "https://www.justdial.com/Mumbai/Hostels-For-Women/nct-10253736"
    req_body = {
        "url": target_url,
        "max_pages": 1,
        "enrich_details": True
    }
    resp = client.post("/api/scrape", json=req_body)
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
    data = resp.json()
    job_id = data.get("job_id")
    assert job_id, "No job_id returned"
    print(f" -> Scrape job started with ID: {job_id}")

    print(" -> Waiting for direct scrape job completion...")
    for _ in range(40):
        time.sleep(1)
        job = jobs.get(job_id)
        if job and job.get("status") in ["completed", "failed"]:
            break

    job = jobs.get(job_id)
    assert job, "Job record missing"
    assert job.get("status") == "completed", f"Job failed or incomplete: {job.get('progress')}"
    results = job.get("results", [])
    print(f" -> Direct scrape completed! Extracted {len(results)} listings.")
    assert len(results) >= 5, f"Expected at least 5 listings, got {len(results)}"

    # ----------------------------------------------------
    # Test 2: Manual Web Search (Top Websites)
    # ----------------------------------------------------
    print("\n[3] Testing POST /api/manual-search (Top Websites Search)...")
    manual_body = {
        "query": "catering services in mumbai",
        "num_websites": 3,
        "enrich_details": True
    }
    resp_manual = client.post("/api/manual-search", json=manual_body)
    assert resp_manual.status_code == 200, f"Expected 200, got {resp_manual.status_code}"
    m_data = resp_manual.json()
    m_job_id = m_data.get("job_id")
    assert m_job_id, "No job_id returned for manual search"
    print(f" -> Manual search job started with ID: {m_job_id}")

    print(" -> Waiting for manual search completion...")
    for _ in range(40):
        time.sleep(1)
        m_job = jobs.get(m_job_id)
        if m_job and m_job.get("status") in ["completed", "failed"]:
            break

    m_job = jobs.get(m_job_id)
    assert m_job, "Manual job record missing"
    assert m_job.get("status") == "completed", f"Manual job failed: {m_job.get('progress')}"
    m_results = m_job.get("results", [])
    print(f" -> Manual search completed! Extracted {len(m_results)} listings across top websites.")
    assert len(m_results) >= 3, f"Expected at least 3 listings from top websites, got {len(m_results)}"

    # ----------------------------------------------------
    # Test 3: Download and Excel Format Verification
    # ----------------------------------------------------
    print("\n[4] Testing GET /api/download/{m_job_id} (Excel download)...")
    resp_dl = client.get(f"/api/download/{m_job_id}")
    assert resp_dl.status_code == 200
    assert "application/vnd.openxmlformats" in resp_dl.headers.get("content-type", "")
    assert len(resp_dl.content) > 1000

    test_download_path = "d:/scraper/downloads/download_verification.xlsx"
    with open(test_download_path, "wb") as f:
        f.write(resp_dl.content)
    print(f" -> Downloaded Excel file ({len(resp_dl.content)} bytes).")

    wb = openpyxl.load_workbook(test_download_path)
    ws = wb.active
    print(f" -> Sheet Name: {ws.title}, Rows: {ws.max_row}, Columns: {ws.max_column}")
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f" -> Excel Headers: {headers}")

    assert "Business Name" in headers
    assert "Phone Number" in headers
    assert "Email ID" in headers
    assert "Full Address" in headers
    assert "Source URL" in headers

    # Cleanup verification file
    if os.path.exists(test_download_path):
        os.remove(test_download_path)

    print("\n" + "=" * 55)
    print("   ALL TESTS (JUSTDIAL & MANUAL SEARCH) PASSED 100%!")
    print("=" * 55)

if __name__ == "__main__":
    test_full_pipeline()
